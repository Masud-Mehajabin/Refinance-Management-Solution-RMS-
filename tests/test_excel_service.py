"""
Tests for excel_service — Excel export functionality.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from services.excel_service import generate_excel, get_excel_filename


class TestGenerateExcel:
    def test_single_record(self, sample_df):
        xlsx = generate_excel(sample_df, branch_id="0101")
        wb = load_workbook(io.BytesIO(xlsx))
        ws = wb.active
        # Header row + 1 data row = 2
        assert ws.max_row == 2
        assert ws.max_column == len(sample_df.columns)

    def test_multiple_records(self, multi_row_df):
        xlsx = generate_excel(multi_row_df, branch_id="0101")
        wb = load_workbook(io.BytesIO(xlsx))
        ws = wb.active
        assert ws.max_row == 4  # header + 3 data rows

    def test_column_names_preserved(self, sample_df):
        xlsx = generate_excel(sample_df)
        wb = load_workbook(io.BytesIO(xlsx))
        ws = wb.active
        header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert header == list(sample_df.columns)

    def test_all_rows_preserved(self, multi_row_df):
        xlsx = generate_excel(multi_row_df)
        wb = load_workbook(io.BytesIO(xlsx))
        ws = wb.active
        data_rows = ws.max_row - 1
        assert data_rows == len(multi_row_df)

    def test_frozen_panes(self, sample_df):
        xlsx = generate_excel(sample_df)
        wb = load_workbook(io.BytesIO(xlsx))
        ws = wb.active
        assert ws.freeze_panes == "A2"

    def test_auto_filter(self, sample_df):
        xlsx = generate_excel(sample_df)
        wb = load_workbook(io.BytesIO(xlsx))
        ws = wb.active
        assert ws.auto_filter.ref is not None

    def test_header_styling(self, sample_df):
        xlsx = generate_excel(sample_df)
        wb = load_workbook(io.BytesIO(xlsx))
        ws = wb.active
        cell = ws.cell(row=1, column=1)
        assert cell.font.bold is True

    def test_empty_dataframe_raises(self, empty_df):
        with pytest.raises(ValueError, match="No data"):
            generate_excel(empty_df)

    def test_none_dataframe_raises(self):
        with pytest.raises(ValueError, match="No data"):
            generate_excel(None)

    def test_unicode_content(self, unicode_df):
        """Unicode data must survive the round-trip."""
        xlsx = generate_excel(unicode_df)
        wb = load_workbook(io.BytesIO(xlsx))
        ws = wb.active
        # Find the BORROWER_NAME column
        header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        name_col = header.index("BORROWER_NAME") + 1
        assert ws.cell(row=2, column=name_col).value == "মোহাম্মদ আবদুল্লাহ"

    def test_output_is_bytes(self, sample_df):
        result = generate_excel(sample_df)
        assert isinstance(result, bytes)
        # Check XLSX magic bytes (PK zip header)
        assert result[:2] == b"PK"


class TestGetExcelFilename:
    def test_contains_branch(self):
        fname = get_excel_filename("0101")
        assert "0101" in fname
        assert fname.endswith(".xlsx")

    def test_contains_timestamp(self):
        fname = get_excel_filename("0101")
        # Should contain YYYYMMDD pattern
        assert any(c.isdigit() for c in fname)

    def test_no_branch(self):
        fname = get_excel_filename()
        assert "Refinance_Details" in fname
        assert fname.endswith(".xlsx")
