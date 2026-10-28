"""HTTP contract tests for the FastAPI application."""

from __future__ import annotations

import json
import sys
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock
from xml.etree.ElementTree import fromstring

import pytest
import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from database import DatabaseError

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import app as app_module
from services.refinance_service import (
    BatchNumberMissingError,
    SubmissionOutcomeUncertainError,
    SubmissionResult,
)


@pytest.fixture
def api_client():
    with TestClient(app_module.app) as client:
        yield client


def test_health_is_available_without_database():
    with TestClient(app_module.app) as client:
        response = client.get("/health")

    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store, no-cache, must-revalidate, max-age=0"
    assert response.headers["pragma"] == "no-cache"
    assert response.json()["status"] == "ok"


def test_dashboard_is_available_without_login():
    with TestClient(app_module.app) as client:
        response = client.get("/")
        logo_response = client.get("/static/meghna-bank-logo.png")

    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store, no-cache, must-revalidate, max-age=0"
    assert "Refinance records" in response.text
    assert 'src="/static/meghna-bank-logo.png"' in response.text
    assert "alt=\"Meghna Bank" in response.text
    assert logo_response.status_code == 200
    assert logo_response.headers["content-type"] == "image/png"
    assert logo_response.headers["cache-control"] == "no-store, no-cache, must-revalidate, max-age=0"
    assert 'id="sidebar-toggle"' in response.text
    assert response.text.index('class="intro-heading"') < response.text.index('id="sidebar-toggle"') < response.text.index("<h1>Refinance records</h1>")
    assert response.text.index('id="sidebar-toggle"') > response.text.index("</header>")
    assert 'aria-expanded="false"' in response.text
    assert 'id="action-sidebar" aria-hidden="true" inert' in response.text
    assert 'id="tag-batch-button"' in response.text
    assert '<span class="nav-label">Tag Reference No.</span>' in response.text
    assert 'class="nav-icon" aria-hidden="true"' in response.text
    assert response.text.index('id="pending-batches-button"') > response.text.index('id="tag-batch-button"')
    assert 'id="pending-batch-number"' in response.text
    assert 'id="retrieve-pending-batches-button"' in response.text
    assert 'id="tag-reference-form"' in response.text
    assert '<select id="tag-reference-batch" required' in response.text
    assert 'id="tag-reference-number" required' in response.text
    assert 'id="cancel-tag-reference"' in response.text
    assert 'id="submit-tag-reference"' in response.text
    assert response.text.index('id="view-batch-details-button"') > response.text.index('id="pending-batches-button"')
    assert 'id="view-batch-details-modal"' in response.text
    assert 'id="cancel-view-batch-details-button"' in response.text
    assert 'id="db-button"' not in response.text
    assert 'id="submit-button"' in response.text
    assert 'id="submit-button" disabled>Submit</button>' in response.text
    assert "Submit / OK" not in response.text
    assert 'id="submit-confirm-modal"' in response.text
    assert 'id="submission-result-modal"' in response.text
    assert 'id="submit-authorizer-username"' not in response.text
    assert 'id="submit-authorizer-password"' not in response.text
    assert "Confirm refinance submission" not in response.text
    assert "Are you sure you want to submit the selected ${selectedIds.length} refinance record(s) for processing?" in response.text
    assert 'id="submit-xml-data" readonly' in response.text
    assert "pmake_by:byId" in response.text
    assert "pauthorize_by:byId" in response.text
    assert 'id="system-dot"' in response.text
    assert 'api("/api/health/database")' in response.text
    assert response.text.index('class="selection-bar"') < response.text.index('class="search"')
    assert 'class="records-toolbar"' in response.text


def test_retrieve_branch_returns_records(api_client, sample_df, monkeypatch):
    monkeypatch.setattr(app_module, "get_refinance_details", lambda branch: sample_df)

    response = api_client.get("/api/refinance/0101")

    assert response.status_code == 200
    payload = response.json()
    assert payload["branch_id"] == "0101"
    assert payload["record_count"] == 1
    assert payload["records"][0]["ACCOUNT_NUMBER"] == "0101000001"
    assert payload["missing_xml_columns"] == []
    loaded_xml = fromstring(payload["xml_data"])
    assert loaded_xml.tag == "TRANSACTION_INFO"
    assert loaded_xml.find("TRANS_DETAILS/Paccount_number").text == "0101000001"
    assert payload["xml_error"] is None


def test_retrieve_all_branches_when_filter_is_omitted(api_client, sample_df, monkeypatch):
    retrieve = Mock(return_value=sample_df)
    monkeypatch.setattr(app_module, "get_refinance_details", retrieve)

    response = api_client.get("/api/refinance")

    assert response.status_code == 200
    assert response.json()["branch_id"] == "ALL BRANCHES"
    retrieve.assert_called_once_with(None)


def _load_snapshot(client, monkeypatch, dataframe):
    monkeypatch.setattr(app_module, "get_refinance_details", lambda branch: dataframe)
    response = client.get("/api/refinance?branch_id=0101")
    assert response.status_code == 200
    return response.json()


def _submission_payload(client, snapshot, selected_ids, **overrides):
    xml_response = client.post(
        "/api/refinance/selected/xml",
        json={"snapshot_id": snapshot["snapshot_id"], "selected_ids": selected_ids},
    )
    assert xml_response.status_code == 200
    return {
        "snapshot_id": snapshot["snapshot_id"],
        "selected_ids": selected_ids,
        "confirmed": True,
        "pxml_data": xml_response.text,
        **overrides,
    }


def test_selected_excel_contains_only_original_selected_rows(api_client, multi_row_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, multi_row_df)
    selected = [snapshot["records"][0]["record_id"], snapshot["records"][2]["record_id"]]
    retrieve = Mock(side_effect=AssertionError("selection export must not query Oracle"))
    monkeypatch.setattr(app_module, "get_refinance_details", retrieve)

    response = api_client.post(
        "/api/refinance/selected/excel",
        json={"snapshot_id": snapshot["snapshot_id"], "selected_ids": selected},
    )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    sheet = workbook.active
    assert sheet.max_row == 3
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert "record_id" not in headers
    account_col = headers.index("ACCOUNT_NUMBER")
    accounts = [row[account_col] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert accounts == [multi_row_df.iloc[0]["ACCOUNT_NUMBER"], multi_row_df.iloc[2]["ACCOUNT_NUMBER"]]
    assert "Refinance_Selected_Records_" in response.headers["content-disposition"]
    retrieve.assert_not_called()


def test_selected_xml_contains_only_original_selected_rows(api_client, multi_row_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, multi_row_df)
    selected = [snapshot["records"][1]["record_id"]]

    response = api_client.post(
        "/api/refinance/selected/xml",
        json={"snapshot_id": snapshot["snapshot_id"], "selected_ids": selected},
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/xml")
    assert "attachment;" in response.headers["content-disposition"]
    assert "Refinance_Selected_Records_" in response.headers["content-disposition"]
    root = fromstring(response.text)
    details = root.findall("TRANS_DETAILS")
    assert len(details) == 1
    assert details[0].get("record_number") == "1"
    assert details[0].findtext("Paccount_number") == multi_row_df.iloc[1]["ACCOUNT_NUMBER"]
    assert all(child.tag != "record_id" for child in details[0])


def test_selected_operations_reject_ids_outside_snapshot(api_client, sample_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, sample_df)

    response = api_client.post(
        "/api/refinance/selected/excel",
        json={"snapshot_id": snapshot["snapshot_id"], "selected_ids": ["not-in-this-snapshot"]},
    )

    assert response.status_code == 422


def test_empty_selection_is_rejected_before_submission(api_client, monkeypatch):
    submission = Mock()
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)

    response = api_client.post(
        "/api/refinance/selected/submit",
        json={
            "snapshot_id": "snapshot",
            "selected_ids": [],
            "confirmed": True,
        },
    )

    assert response.status_code == 422
    submission.assert_not_called()


def test_selected_batch_submission_requires_enablement(api_client, sample_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, sample_df)
    submission = Mock()
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", False)
    payload = _submission_payload(api_client, snapshot, [snapshot["records"][0]["record_id"]])

    response = api_client.post(
        "/api/refinance/selected/submit",
        json=payload,
    )

    assert response.status_code == 403
    submission.assert_not_called()


def test_selected_batch_requires_xml_data(api_client, sample_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, sample_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    submission = Mock()
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)

    response = api_client.post(
        "/api/refinance/selected/submit",
        json={
            "snapshot_id": snapshot["snapshot_id"],
            "selected_ids": [snapshot["records"][0]["record_id"]],
            "confirmed": True,
        },
    )

    assert response.status_code == 422
    submission.assert_not_called()


@pytest.mark.parametrize("invalid_xml", ["<TRANSACTION_INFO>", "   "])
def test_selected_batch_rejects_invalid_xml(api_client, sample_df, monkeypatch, invalid_xml):
    snapshot = _load_snapshot(api_client, monkeypatch, sample_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    submission = Mock()
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)
    payload = _submission_payload(
        api_client,
        snapshot,
        [snapshot["records"][0]["record_id"]],
        pxml_data=invalid_xml,
    )

    response = api_client.post("/api/refinance/selected/submit", json=payload)

    assert response.status_code == 422
    submission.assert_not_called()


def test_selected_batch_rejects_xml_that_does_not_match_selection(api_client, multi_row_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, multi_row_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    submission = Mock()
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)
    selected_ids = [snapshot["records"][0]["record_id"]]
    different_xml_response = api_client.post(
        "/api/refinance/selected/xml",
        json={
            "snapshot_id": snapshot["snapshot_id"],
            "selected_ids": [snapshot["records"][1]["record_id"]],
        },
    )
    payload = _submission_payload(
        api_client,
        snapshot,
        selected_ids,
        pxml_data=different_xml_response.text,
    )

    response = api_client.post("/api/refinance/selected/submit", json=payload)

    assert response.status_code == 422
    assert "does not match" in response.json()["detail"]
    submission.assert_not_called()


def test_selected_batch_submission_passes_optional_parameters_as_null(api_client, sample_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, sample_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    submission = Mock()
    submission.return_value = SubmissionResult(
        success=True,
        message="000 - Successfully inserted 1 record(s)",
        batch_number="B-001",
        record_count=1,
    )
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)

    payload = _submission_payload(api_client, snapshot, [snapshot["records"][0]["record_id"]])
    response = api_client.post(
        "/api/refinance/selected/submit",
        json=payload,
    )

    assert response.status_code == 200
    assert submission.call_args.kwargs["make_by"] is None
    assert submission.call_args.kwargs["make_date"] is None
    assert submission.call_args.kwargs["authorize_by"] is None
    assert submission.call_args.kwargs["authorize_date"] is None
    assert submission.call_args.kwargs["xml_data"] == payload["pxml_data"]


def test_selected_batch_submission_requires_explicit_confirmation(api_client, sample_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, sample_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    submission = Mock()
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)
    payload = _submission_payload(
        api_client,
        snapshot,
        [snapshot["records"][0]["record_id"]],
        confirmed=False,
    )

    response = api_client.post(
        "/api/refinance/selected/submit",
        json=payload,
    )

    assert response.status_code == 400
    submission.assert_not_called()


def test_selected_batch_submission_uses_only_selected_original_rows(api_client, multi_row_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, multi_row_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    result = SubmissionResult(
        success=True,
        message="000 - Successfully inserted 2 record(s)",
        batch_number="BATCH-001",
        record_count=2,
    )
    submission = Mock(return_value=result)
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)

    payload = _submission_payload(
        api_client,
        snapshot,
        [snapshot["records"][0]["record_id"], snapshot["records"][2]["record_id"]],
        pmake_by="maker-user",
        pmake_date="2024-03-15",
        pauthorize_by="authorizer-user",
        pauthorize_date="2024-03-16",
    )
    response = api_client.post(
        "/api/refinance/selected/submit",
        json=payload,
    )

    assert response.status_code == 200
    assert response.json()["batch_number"] == "BATCH-001"
    assert response.json()["record_count"] == 2
    assert submission.call_args.kwargs["make_by"] == "maker-user"
    assert submission.call_args.kwargs["make_date"].isoformat() == "2024-03-15T00:00:00"
    assert submission.call_args.kwargs["authorize_by"] == "authorizer-user"
    assert submission.call_args.kwargs["authorize_date"].isoformat() == "2024-03-16T00:00:00"
    assert submission.call_args.kwargs["xml_data"] == payload["pxml_data"]
    xml_root = fromstring(submission.call_args.kwargs["xml_data"])
    assert len(xml_root.findall("TRANS_DETAILS")) == 2
    assert [detail.get("record_number") for detail in xml_root.findall("TRANS_DETAILS")] == ["1", "2"]
    assert [detail.findtext("Paccount_number") for detail in xml_root.findall("TRANS_DETAILS")] == [
        multi_row_df.iloc[0]["ACCOUNT_NUMBER"],
        multi_row_df.iloc[2]["ACCOUNT_NUMBER"],
    ]


def test_selected_batch_blocks_empty_required_source_fields(api_client, sample_df, monkeypatch):
    incomplete_df = sample_df.drop(columns=["ACCOUNT_TITLE"])
    snapshot = _load_snapshot(api_client, monkeypatch, incomplete_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    submission = Mock(
        return_value=SubmissionResult(
            success=True,
            message="000 - Successfully inserted 1 record(s)",
            batch_number="B-001",
            record_count=1,
        )
    )
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)

    response = api_client.post(
        "/api/refinance/selected/submit",
        json=_submission_payload(api_client, snapshot, [snapshot["records"][0]["record_id"]]),
    )

    assert response.status_code == 422
    assert "Paccount_title" in response.json()["detail"]
    submission.assert_not_called()


def test_selected_batch_does_not_report_success_without_procedure_confirmation(
    api_client, sample_df, monkeypatch
):
    snapshot = _load_snapshot(api_client, monkeypatch, sample_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    monkeypatch.setattr(
        app_module,
        "submit_refinance_batch",
        Mock(return_value=SubmissionResult(success=False, message="No batch number returned")),
    )

    response = api_client.post(
        "/api/refinance/selected/submit",
        json=_submission_payload(api_client, snapshot, [snapshot["records"][0]["record_id"]]),
    )

    assert response.status_code == 502
    assert response.json()["detail"] == "No batch number returned"


def test_selected_batch_rejects_duplicate_submission(api_client, sample_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, sample_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    submission = Mock(return_value=SubmissionResult(success=True, message="ok", batch_number="B-1", record_count=1))
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)
    payload = _submission_payload(api_client, snapshot, [snapshot["records"][0]["record_id"]])

    first = api_client.post("/api/refinance/selected/submit", json=payload)
    second = api_client.post("/api/refinance/selected/submit", json=payload)

    assert first.status_code == 200
    assert second.status_code == 409
    submission.assert_called_once()


def test_selected_batch_can_retry_after_confirmed_database_rollback(api_client, sample_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, sample_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    submission = Mock(
        side_effect=[
            DatabaseError("Oracle rejected input; transaction rolled back."),
            SubmissionResult(
                success=True,
                message="000 - Successfully inserted 1 record(s)",
                batch_number="B-RETRY",
                record_count=1,
            ),
        ]
    )
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)
    payload = _submission_payload(api_client, snapshot, [snapshot["records"][0]["record_id"]])

    first = api_client.post("/api/refinance/selected/submit", json=payload)
    retry = api_client.post("/api/refinance/selected/submit", json=payload)

    assert first.status_code == 503
    assert retry.status_code == 200
    assert retry.json()["batch_number"] == "B-RETRY"
    assert submission.call_count == 2


def test_selected_batch_blocks_retry_after_uncertain_commit(api_client, sample_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, sample_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    submission = Mock(
        side_effect=SubmissionOutcomeUncertainError(
            "The database commit outcome is uncertain. Verify the batch status before retrying."
        )
    )
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)
    payload = _submission_payload(api_client, snapshot, [snapshot["records"][0]["record_id"]])

    first = api_client.post("/api/refinance/selected/submit", json=payload)
    retry = api_client.post("/api/refinance/selected/submit", json=payload)

    assert first.status_code == 503
    assert retry.status_code == 409
    submission.assert_called_once()


def test_selected_batch_rejects_overlapping_submission(api_client, multi_row_df, monkeypatch):
    snapshot = _load_snapshot(api_client, monkeypatch, multi_row_df)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    submission = Mock(
        return_value=SubmissionResult(
            success=True,
            message="Inserted",
            batch_number="B-2",
            record_count=2,
        )
    )
    monkeypatch.setattr(app_module, "submit_refinance_batch", submission)
    record_ids = [record["record_id"] for record in snapshot["records"]]

    def request_for(ids):
        return _submission_payload(api_client, snapshot, ids)

    first = api_client.post(
        "/api/refinance/selected/submit",
        json=request_for(record_ids[:2]),
    )
    overlapping = api_client.post(
        "/api/refinance/selected/submit",
        json=request_for(record_ids[1:]),
    )

    assert first.status_code == 200
    assert overlapping.status_code == 409
    submission.assert_called_once()

def _tag_batch_payload():
    return {}


def test_tag_batch_requires_write_enablement(api_client, monkeypatch):
    generate = Mock()
    monkeypatch.setattr(app_module, "generate_batch_number", generate)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", False)

    response = api_client.post("/api/refinance/tag-batch", json=_tag_batch_payload())

    assert response.status_code == 403
    generate.assert_not_called()


@pytest.mark.parametrize("batch_number", ["12", "BATCH-87"])
def test_tag_batch_returns_actual_number_and_passes_null_optional_inputs(
    api_client, monkeypatch, batch_number
):
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    generate = Mock(return_value=batch_number)
    monkeypatch.setattr(app_module, "generate_batch_number", generate)

    response = api_client.post("/api/refinance/tag-batch", json=_tag_batch_payload())

    assert response.status_code == 200
    assert response.json() == {
        "success": True,
        "batch_number": batch_number,
        "message": f"Batch {batch_number} has been generated successfully.",
    }
    generate.assert_called_once_with(total_loan_no=None, create_by=None)


def test_tag_batch_passes_entered_optional_inputs(api_client, monkeypatch):
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    generate = Mock(return_value="88")
    monkeypatch.setattr(app_module, "generate_batch_number", generate)
    payload = {
        "total_loan_no": "125",
        "create_by": "maker-user",
    }

    response = api_client.post("/api/refinance/tag-batch", json=payload)

    assert response.status_code == 200
    assert response.json()["batch_number"] == "88"
    generate.assert_called_once_with(total_loan_no="125", create_by="maker-user")


def test_tag_batch_does_not_report_success_without_output(api_client, monkeypatch):
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    generate = Mock(
        side_effect=BatchNumberMissingError(
            "Batch generation completed, but no batch number was returned. "
            "Please verify the database procedure."
        )
    )
    monkeypatch.setattr(app_module, "generate_batch_number", generate)

    response = api_client.post("/api/refinance/tag-batch", json=_tag_batch_payload())

    assert response.status_code == 502
    assert response.json()["detail"].startswith("Batch generation completed")


def test_tag_batch_reports_oracle_error_safely(api_client, monkeypatch):
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    generate = Mock(
        side_effect=app_module.DatabaseError(
            "Unable to generate the batch. Please contact the system administrator if the problem persists."
        )
    )
    monkeypatch.setattr(app_module, "generate_batch_number", generate)

    response = api_client.post("/api/refinance/tag-batch", json=_tag_batch_payload())

    assert response.status_code == 503
    assert "system administrator" in response.json()["detail"]


def test_tag_batch_rejects_concurrent_repeat(api_client, monkeypatch):
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    generate = Mock(return_value="12")
    monkeypatch.setattr(app_module, "generate_batch_number", generate)
    assert app_module._tag_batch_lock.acquire(blocking=False)
    try:
        response = api_client.post("/api/refinance/tag-batch", json=_tag_batch_payload())
    finally:
        app_module._tag_batch_lock.release()

    assert response.status_code == 409
    generate.assert_not_called()


def test_tag_reference_requires_write_enablement(api_client, monkeypatch):
    update = Mock()
    monkeypatch.setattr(app_module, "tag_batch_reference", update)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", False)

    response = api_client.post(
        "/api/refinance/tag-reference",
        json={"batch_no": "12", "reference_no": "BB-REF-01"},
    )

    assert response.status_code == 403
    update.assert_not_called()


def test_tag_reference_batch_options_returns_batches_from_service(api_client, monkeypatch):
    batches = ["10", "11", "12"]
    retrieve = Mock(return_value=batches)
    monkeypatch.setattr(app_module, "get_refinance_batch_numbers", retrieve)

    response = api_client.get("/api/refinance/tag-reference/batches")

    assert response.status_code == 200
    assert response.json() == {"batch_numbers": batches}
    retrieve.assert_called_once_with()


def test_tag_reference_batch_options_reports_database_error_safely(api_client, monkeypatch):
    retrieve = Mock(
        side_effect=app_module.DatabaseError(
            "Unable to retrieve batch numbers. Please try again or contact the system administrator."
        )
    )
    monkeypatch.setattr(app_module, "get_refinance_batch_numbers", retrieve)

    response = api_client.get("/api/refinance/tag-reference/batches")

    assert response.status_code == 503
    assert "Unable to retrieve batch numbers" in response.json()["detail"]


@pytest.mark.parametrize(
    ("payload", "message"),
    [
        ({"batch_no": " ", "reference_no": "BB-REF"}, "Batch No. is required."),
        ({"batch_no": "12", "reference_no": "  "}, "Reference No. is required."),
    ],
)
def test_tag_reference_rejects_blank_values(api_client, monkeypatch, payload, message):
    update = Mock()
    monkeypatch.setattr(app_module, "tag_batch_reference", update)
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)

    response = api_client.post("/api/refinance/tag-reference", json=payload)

    assert response.status_code == 422
    assert response.json()["detail"] == message
    update.assert_not_called()


def test_tag_reference_trims_values_and_returns_actual_update_count(api_client, monkeypatch):
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    update = Mock(return_value=12)
    monkeypatch.setattr(app_module, "tag_batch_reference", update)

    response = api_client.post(
        "/api/refinance/tag-reference",
        json={"batch_no": "  0012 ", "reference_no": " BB-REF-2026-001  "},
    )

    assert response.status_code == 200
    assert response.json() == {
        "success": True,
        "batch_no": "0012",
        "reference_no": "BB-REF-2026-001",
        "records_updated": 12,
        "message": "Reference No. BB-REF-2026-001 has been successfully assigned to "
        "Batch No. 0012. Total records updated: 12.",
    }
    update.assert_called_once_with("0012", "BB-REF-2026-001")


def test_tag_reference_returns_not_found_without_success(api_client, monkeypatch):
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    monkeypatch.setattr(app_module, "tag_batch_reference", Mock(return_value=0))

    response = api_client.post(
        "/api/refinance/tag-reference",
        json={"batch_no": "404", "reference_no": "BB-REF"},
    )

    assert response.status_code == 404
    assert response.json()["detail"] == (
        "No records found for the specified Batch No. Please verify the batch number and try again."
    )
    assert response.json().get("success") is None


def test_tag_reference_reports_database_error_safely(api_client, monkeypatch):
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    update = Mock(
        side_effect=app_module.DatabaseError(
            "Unable to assign the reference number due to a database error. "
            "Please try again or contact the system administrator."
        )
    )
    monkeypatch.setattr(app_module, "tag_batch_reference", update)

    response = api_client.post(
        "/api/refinance/tag-reference",
        json={"batch_no": "12", "reference_no": "BB-REF"},
    )

    assert response.status_code == 503
    assert "Unable to assign the reference number" in response.json()["detail"]


def test_tag_reference_rejects_concurrent_repeat(api_client, monkeypatch):
    monkeypatch.setattr(app_module, "ALLOW_BATCH_SUBMISSION", True)
    update = Mock()
    monkeypatch.setattr(app_module, "tag_batch_reference", update)
    assert app_module._tag_reference_lock.acquire(blocking=False)
    try:
        response = api_client.post(
            "/api/refinance/tag-reference",
            json={"batch_no": "12", "reference_no": "BB-REF"},
        )
    finally:
        app_module._tag_reference_lock.release()

    assert response.status_code == 409
    update.assert_not_called()


def test_pending_batches_passes_optional_batch_number_and_does_not_tag(api_client, monkeypatch):
    results = pd.DataFrame(
        [{"BATCH_NO": "0012", "TOTAL_LOAN_AMOUNT": 1250, "NO_OF_LOAN": 2}]
    )
    retrieve = Mock(return_value=results)
    generate = Mock()
    monkeypatch.setattr(app_module, "get_pending_batches", retrieve)
    monkeypatch.setattr(app_module, "generate_batch_number", generate)

    response = api_client.post(
        "/api/refinance/pending-batches",
        json={"batch_no": "  0012  "},
    )

    assert response.status_code == 200
    assert response.json() == {
        "batch_no": "0012",
        "columns": ["BATCH_NO", "TOTAL_LOAN_AMOUNT", "NO_OF_LOAN"],
        "records": [{"BATCH_NO": "0012", "TOTAL_LOAN_AMOUNT": 1250, "NO_OF_LOAN": 2}],
        "record_count": 1,
    }
    retrieve.assert_called_once_with("0012")
    generate.assert_not_called()


def test_pending_batches_blank_filter_passes_null_and_empty_results(api_client, monkeypatch):
    retrieve = Mock(return_value=pd.DataFrame(columns=["BATCH_NO", "MAKE_BY"]))
    generate = Mock()
    monkeypatch.setattr(app_module, "get_pending_batches", retrieve)
    monkeypatch.setattr(app_module, "generate_batch_number", generate)

    response = api_client.post("/api/refinance/pending-batches", json={"batch_no": "  "})

    assert response.status_code == 200
    assert response.json() == {
        "batch_no": None,
        "columns": ["BATCH_NO", "MAKE_BY"],
        "records": [],
        "record_count": 0,
    }
    retrieve.assert_called_once_with(None)
    generate.assert_not_called()


def test_pending_batches_reports_safe_database_error(api_client, monkeypatch):
    retrieve = Mock(side_effect=app_module.DatabaseError("Unable to retrieve pending batches."))
    monkeypatch.setattr(app_module, "get_pending_batches", retrieve)

    response = api_client.post("/api/refinance/pending-batches", json={"batch_no": None})

    assert response.status_code == 503
    assert "Unable to retrieve pending batches" in response.json()["detail"]


def test_batch_details_passes_optional_value_and_does_not_run_other_batch_procedures(
    api_client, monkeypatch
):
    results = pd.DataFrame(
        [{"BATCH_NO": "0012", "LOAN_ACCOUNT_NUMBER": "00001234"}]
    )
    retrieve = Mock(return_value=results)
    generate = Mock()
    pending = Mock()
    monkeypatch.setattr(app_module, "get_batch_details", retrieve)
    monkeypatch.setattr(app_module, "generate_batch_number", generate)
    monkeypatch.setattr(app_module, "get_pending_batches", pending)

    response = api_client.post(
        "/api/refinance/batch-details",
        json={"batch_no": "  0012  "},
    )

    assert response.status_code == 200
    assert response.json() == {
        "batch_no": "0012",
        "columns": ["BATCH_NO", "LOAN_ACCOUNT_NUMBER"],
        "records": [{"BATCH_NO": "0012", "LOAN_ACCOUNT_NUMBER": "00001234"}],
        "record_count": 1,
    }
    retrieve.assert_called_once_with("0012")
    generate.assert_not_called()
    pending.assert_not_called()


def test_batch_details_blank_filter_and_empty_results(api_client, monkeypatch):
    retrieve = Mock(return_value=pd.DataFrame(columns=["BATCH_NO", "MAKE_BY"]))
    monkeypatch.setattr(app_module, "get_batch_details", retrieve)

    response = api_client.post("/api/refinance/batch-details", json={"batch_no": "  "})

    assert response.status_code == 200
    assert response.json() == {
        "batch_no": None,
        "columns": ["BATCH_NO", "MAKE_BY"],
        "records": [],
        "record_count": 0,
    }
    retrieve.assert_called_once_with(None)


def test_batch_details_reports_safe_database_error(api_client, monkeypatch):
    retrieve = Mock(side_effect=app_module.DatabaseError("Unable to retrieve batch details."))
    monkeypatch.setattr(app_module, "get_batch_details", retrieve)

    response = api_client.post("/api/refinance/batch-details", json={"batch_no": None})

    assert response.status_code == 503
    assert "Unable to retrieve batch details" in response.json()["detail"]