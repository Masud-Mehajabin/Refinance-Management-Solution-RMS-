"""
Excel Service — exports refinance data to formatted .xlsx workbooks.

Uses pandas + openpyxl for writing, then applies header formatting,
frozen panes, and auto-filters.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("refinance_app.services.excel_service")


def generate_excel(df: pd.DataFrame, branch_id: str = "") -> bytes:
    """
    Export *df* to an in-memory .xlsx file and return the raw bytes.

    Parameters
    ----------
    df : pd.DataFrame
        The refinance records to export.
    branch_id : str
        Used only for the sheet name / metadata.

    Returns
    -------
    bytes
        The complete .xlsx file as a byte string, ready for download.

    Raises
    ------
    ValueError
        If the DataFrame is empty.
    """
    if df is None or df.empty:
        raise ValueError("No data available for Excel export.")

    buf = io.BytesIO()

    sheet_name = f"Refinance_{branch_id}" if branch_id else "Refinance_Data"
    # Sheet names max 31 chars
    sheet_name = sheet_name[:31]

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        wb = writer.book
        ws = writer.sheets[sheet_name]

        # --- Header styling ---------------------------------------------------
        header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="1B3A5C", end_color="1B3A5C", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # --- Data cell formatting ---------------------------------------------
        data_font = Font(name="Calibri", size=10)
        data_align = Alignment(vertical="center", wrap_text=False)

        for row_idx in range(2, len(df) + 2):
            for col_idx in range(1, len(df.columns) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.alignment = data_align
                cell.border = thin_border

        # --- Column widths (auto-fit approximation) ---------------------------
        for col_idx, col_name in enumerate(df.columns, start=1):
            max_len = len(str(col_name))
            for row_idx in range(2, min(len(df) + 2, 102)):  # Sample first 100 rows
                cell_val = ws.cell(row=row_idx, column=col_idx).value
                if cell_val is not None:
                    max_len = max(max_len, len(str(cell_val)))
            # Clamp between 10 and 40
            adjusted = min(max(max_len + 2, 10), 40)
            ws.column_dimensions[get_column_letter(col_idx)].width = adjusted

        # --- Freeze header row ------------------------------------------------
        ws.freeze_panes = "A2"

        # --- Auto-filter on header row ----------------------------------------
        last_col = get_column_letter(len(df.columns))
        ws.auto_filter.ref = f"A1:{last_col}{len(df) + 1}"

    buf.seek(0)
    logger.info("Excel file generated: %d rows, %d columns.", len(df), len(df.columns))
    return buf.getvalue()


def get_excel_filename(branch_id: str = "") -> str:
    """Return a descriptive filename with branch and timestamp."""
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    branch_part = f"_{branch_id}" if branch_id else ""
    return f"Refinance_Details{branch_part}_{ts}.xlsx"
